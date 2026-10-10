# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Build the one-day decks from the three session decks.

The three session decks in site/public/slides/ are the masters. The one-day
decks are copies with a few slides hidden, one slide moved and a few lines of
text changed (see site/src/content/facilitator/one-day.md). Never edit the
one-day decks by hand: change a master, then rebuild.

    just one-day-decks          # rebuild
    just one-day-decks --check  # fail if any one-day deck is out of date

Slides are found by their text, not their number, so a master can be
reordered safely. If a master's wording changes so a slide can't be found,
the build stops and names it: update the table below to match.
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SLIDES = ROOT / "site" / "public" / "slides"
STAMP = "one-day build of {src} sha256:{sha}"

BANNER = (
    "ONE-DAY VERSION. Built from {src} by tools/one_day_decks.py: don't edit "
    "this file, edit the master and rebuild. Timings and running order: "
    "/facilitator/one-day. The other notes come from the three-session "
    "version: read \"next week\" as \"after lunch\" or \"in the 1-week email\"."
)

# Each deck: source, output, slides to hide (by a phrase on the slide), text
# swaps as (phrase on the slide, exact run text, new text), each of which must
# match, and optional moves (slide phrase -> put it after this slide phrase).
DECKS = [
    dict(
        src="ttl-day1-death-by-powerpoint.pptx",
        out="ttl-one-day-1-death-by-powerpoint.pptx",
        hide=["The ask", "See you next week"],
        swaps=[
            ("Death by PowerPoint?", "Day 1 of 3", "In one day · 1 of 3"),
            ("Building your session", "Day 1", "This morning"),
            ("Building your session", "Day 2", "Before lunch"),
            ("Building your session", "Day 3", "This afternoon"),
        ],
    ),
    dict(
        src="ttl-day2-active-beats-passive.pptx",
        out="ttl-one-day-2-active-beats-passive.pptx",
        hide=["Show us your slides", "60-second pitch", "The ask",
              "Next week: did it land?", "See you next week"],
        swaps=[
            ("Active Beats Passive?", "Day 2 of 3", "In one day · 2 of 3"),
            ("Building your session", "Day 2: what they do", "Before lunch: what they do"),
            ("Building your session", "Day 1 ✓ ", "This morning ✓ "),
            ("Building your session", "Day 3: ", "This afternoon: "),
            ("Quick quiz:", "last week", "this morning"),
            ("What did you just", "A week later", "Hours later"),
        ],
    ),
    dict(
        src="ttl-day3-did-it-land.pptx",
        out="ttl-one-day-3-did-it-land.pptx",
        hide=[],
        swaps=[
            ("Did It Land?", "Day 3 of 3", "In one day · 3 of 3"),
            ("Building your session", "Days 1–2 ✓", "This morning ✓"),
            ("Building your session", "Day 3: ", "This afternoon: "),
            ("You've been on a canvas", "for ", ""),
            ("You've been on a canvas", "three weeks", "all day"),
            ("Here's ours", "An ask each week", "A 1-week email"),
        ],
        # The blorps retest opens the last block, after the break.
        move={"Who remembers what a blorp is?": "Your verb decides your evidence"},
    ),
]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def slide_text(slide) -> str:
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def find(prs, phrase: str, src: str):
    hits = [s for s in prs.slides if phrase in slide_text(s)]
    if not hits:
        sys.exit(f"{src}: no slide contains {phrase!r}. Update DECKS in tools/one_day_decks.py.")
    return hits[0]


def build(d: dict) -> None:
    src = SLIDES / d["src"]
    prs = Presentation(src)

    for phrase in d["hide"]:
        find(prs, phrase, d["src"])._element.set("show", "0")

    for phrase, old, new in d["swaps"]:
        slide = find(prs, phrase, d["src"])
        runs = [r for sh in slide.shapes if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs if r.text == old]
        if not runs:
            sys.exit(f"{d['src']}: {old!r} not found on the {phrase!r} slide. "
                     "Update DECKS in tools/one_day_decks.py.")
        for r in runs:
            r.text = new

    ids = prs.slides._sldIdLst
    entry = lambda slide: next(e for e in ids if e.id == slide.slide_id)
    for phrase, after in d.get("move", {}).items():
        moving = entry(find(prs, phrase, d["src"]))
        entry(find(prs, after, d["src"])).addnext(moving)

    first = prs.slides[0]
    notes = first.notes_slide.notes_text_frame if first.has_notes_slide else None
    if notes is not None:
        notes.text = BANNER.format(src=d["src"]) + ("\n\n" + notes.text if notes.text else "")

    prs.core_properties.comments = STAMP.format(src=d["src"], sha=sha(src))
    prs.save(SLIDES / d["out"])
    print(f"built {d['out']}")


def stale(d: dict) -> str | None:
    out = SLIDES / d["out"]
    if not out.exists():
        return "missing"
    want = STAMP.format(src=d["src"], sha=sha(SLIDES / d["src"]))
    return None if Presentation(out).core_properties.comments == want else "out of date"


if __name__ == "__main__":
    if "--check" in sys.argv:
        bad = [(d["out"], why) for d in DECKS if (why := stale(d))]
        for out, why in bad:
            print(f"{out}: {why}. Run: just one-day-decks")
        sys.exit(1 if bad else 0)
    for d in DECKS:
        build(d)
